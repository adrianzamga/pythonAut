from openpyxl import load_workbook

archivo = ".xlsx"
wb = load_workbook(archivo)
ws = wb.active

# Detecta la última fila con datos
ultima_fila = ws.max_row

for i in range(ultima_fila, 0, -1):
    # Cada bloque tiene 60 filas:
    # 55 se conservan y 5 se borran
    if (i % 60) > 55 or (i % 60) == 0:
        ws.delete_rows(i, 1)

wb.save("archivo_limpio.xlsx") 